from __future__ import annotations

import io
import re
from datetime import datetime

import httpx
import pandas as pd
from lxml import html
from openpyxl import load_workbook

from ..catalog import TIM_COUNTRY
from ..http import get_bytes, get_text
from .base import run_connector

INDEX = "https://tim.org.tr/tr/ihracat-rakamlari"
SECTOR = "Hububat, Bakliyat, Yağlı Tohumlar ve Mamulleri"


def _latest_xlsx(c: httpx.Client) -> tuple[str, bytes]:
    page = get_text(c, INDEX)
    hrefs = re.findall(r'href="([^"]+ulkelere-gore-sektorel-ihracat-rakamlari\.xlsx[^"]*)"', page)
    if not hrefs:
        raise RuntimeError("TİM sector-by-country workbook not found")
    href = hrefs[0]
    if href.startswith("/"):
        href = "https://tim.org.tr" + href
    return href, get_bytes(c, href)


def parse_workbook(raw: bytes) -> pd.DataFrame:
    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    recs = []
    for row in rows[4:]:
        if not row or len(row) < 5:
            continue
        sector = (row[0] or "")
        country = (row[1] or "")
        if SECTOR not in str(sector):
            continue
        recs.append(
            {
                "sector": str(sector).strip(),
                "country_tr": str(country).strip(),
                "iso2": TIM_COUNTRY.get(str(country).strip()),
                "sep_yoy_prev": row[2],
                "sep_yoy_now": row[3],
                "aug_now": row[5],
                "ytd_prev": row[7],
                "ytd_now": row[8],
                "source": "tim",
            }
        )
    df = pd.DataFrame(recs)
    for col in ("sep_yoy_prev", "sep_yoy_now", "aug_now", "ytd_prev", "ytd_now"):
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    # TİM values are thousands of USD
    for col in ("sep_yoy_prev", "sep_yoy_now", "aug_now", "ytd_prev", "ytd_now"):
        if col in df.columns:
            df[col + "_usd"] = df[col] * 1000
    df["yoy_sep"] = (df["sep_yoy_now"] - df["sep_yoy_prev"]) / df["sep_yoy_prev"].replace(0, pd.NA)
    df["yoy_ytd"] = (df["ytd_now"] - df["ytd_prev"]) / df["ytd_prev"].replace(0, pd.NA)
    return df


def fetch(c: httpx.Client) -> pd.DataFrame:
    href, raw = _latest_xlsx(c)
    df = parse_workbook(raw)
    df["source_url"] = href
    df["as_of"] = datetime.utcnow().strftime("%Y-%m")
    return df


def refresh(c: httpx.Client):
    return run_connector(
        "tim_exports",
        "tim_exports",
        INDEX,
        "Turkish Exporters Assembly monthly figures, public website",
        lambda: fetch(c),
        notes="Cereals/pulses/oilseeds sector exports by destination (includes pasta). Values in USD.",
        as_of_col="as_of",
    )


_ = html  # imported for future listing parse if TİM markup changes
